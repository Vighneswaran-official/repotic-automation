"""
tests/test_all.py - Complete verification test suite for REPOTIC automation.
"""

import sys
from pathlib import Path
from openpyxl import load_workbook

_ROOT = Path(__file__).resolve().parent.parent
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from src.config import fin_year, PREFIX, detect_month_from_repotic, resolve_report_period
from src.services.processor import run_automation


def run_tests():
    print("=" * 80)
    print("STARTING REPOTIC AUTOMATION VERIFICATION SUITE")
    print("=" * 80)

    # ── Test 3: Unit-test fin_year ───────────────────────────────────────────
    print("\n--- Test 3: Unit-testing fin_year ---")
    fy_test_cases = [
        ((3, 2027), "26-27"),
        ((4, 2027), "27-28"),
        ((1, 2027), "26-27"),
        ((12, 2026), "26-27"),
        ((4, 2026), "26-27"),
        ((3, 2026), "25-26"),
    ]
    for (m, y), expected in fy_test_cases:
        actual = fin_year(m, y)
        assert actual == expected, f"fin_year({m}, {y}) returned {actual}, expected {expected}"
        print(f"  ✓ fin_year({m}, {y}) -> {actual}")
    print("Test 3 PASSED!")

    repotic_file  = _ROOT / "output" / "upload_repotic.xlsx"
    ledger_file   = _ROOT / "output" / "upload_ledger.xlsx"
    template_file = _ROOT / "output" / "upload_template.xlsx"
    test_output   = _ROOT / "output" / "Final_Output_Tamilnadu.xlsx"

    # ── Test 5: Confirm month detection ──────────────────────────────────────
    print("\n--- Test 5: Auto-detect report month when no month is selected in UI ---")
    year, month, source = resolve_report_period(repotic_path=repotic_file, month_override=None)
    print(f"  Detected month: {year}-{month:02d} (Source: {source})")
    assert month == 8 and year == 2026, f"Expected August 2026 (8, 2026), got ({month}, {year})"
    print("Test 5 PASSED!")

    # ── Test 1 & 2: Run with current data (August 2026) ─────────────────────
    print("\n--- Test 1 & 2: Execution with current data (August 2026) ---")
    res = run_automation(repotic_file, ledger_file, template_file, test_output, month_override="2026-08")

    # Check series for each marketplace
    series_map = {s["sheet"]: s["invoice_series"] for s in res["summary"]}
    print(f"  Flipkart series: {series_map.get('Flipkart')}")
    print(f"  Meesho series  : {series_map.get('Meesho')}")
    print(f"  Snapdeal series: {series_map.get('Snapdeal')}")

    assert series_map.get("Flipkart") == "08/26-27/FL-01 to 08/26-27/FL-31", f"Flipkart series mismatch: {series_map.get('Flipkart')}"
    assert series_map.get("Meesho") == "08/26-27/ME-01 to 08/26-27/ME-26", f"Meesho series mismatch: {series_map.get('Meesho')}"
    assert series_map.get("Snapdeal") == "08/26-27/SN-01 to 08/26-27/SN-28", f"Snapdeal series mismatch: {series_map.get('Snapdeal')}"
    print("Test 1 PASSED: Flipkart is FL-01 to FL-31, Meesho is ME-01 to ME-26, Snapdeal is SN-01 to SN-28.")

    # Test 2: Uniqueness of InvNo across output
    wb = load_workbook(test_output, data_only=False)
    ws = wb.active
    header_row = [cell.value for cell in ws[1]]
    inv_col = header_row.index("InvNo") + 1

    inv_numbers = [ws.cell(r, inv_col).value for r in range(2, ws.max_row + 1)]
    assert len(inv_numbers) == len(set(inv_numbers)), f"Duplicate invoice numbers found! Total: {len(inv_numbers)}, Unique: {len(set(inv_numbers))}"
    print(f"Test 2 PASSED: All {len(inv_numbers)} invoice numbers are distinct. No InvNo appears twice.")

    # ── Test 6: Confirm InvNo cells are text in Excel ────────────────────────
    print("\n--- Test 6: Confirm InvNo cells are stored as text in Excel ---")
    for r in range(2, min(10, ws.max_row + 1)):
        c = ws.cell(r, inv_col)
        assert isinstance(c.value, str), f"Row {r} InvNo is not string: {type(c.value)}"
        assert c.number_format == "@", f"Row {r} InvNo number format is {c.number_format}, expected '@'"
    print("Test 6 PASSED: InvNo cells have number_format='@' and python type str.")

    # ── Test 4: CLI with --month 2027-04 ─────────────────────────────────────
    print("\n--- Test 4: Execution with --month 2027-04 ---")
    test_output_2027 = _ROOT / "output" / "Test_Output_2027.xlsx"
    res_2027 = run_automation(repotic_file, ledger_file, template_file, test_output_2027, month_override="2027-04")
    series_2027 = {s["sheet"]: s["invoice_series"] for s in res_2027["summary"]}
    print(f"  Flipkart 2027 series: {series_2027.get('Flipkart')}")
    assert series_2027.get("Flipkart").startswith("04/27-28/FL-01"), f"2027 series failed: {series_2027.get('Flipkart')}"
    if test_output_2027.exists():
        test_output_2027.unlink()
    # ── Test 7: Confirm HSNCode is blank across all rows (no HSN needed) ───
    print("\n--- Test 7: Confirm HSNCode is blank across all rows ---")
    assert "HSNCode" in header_row, "HSNCode header missing from template schema"
    hsn_col = header_row.index("HSNCode") + 1
    hsn_values = [ws.cell(r, hsn_col).value for r in range(2, ws.max_row + 1)]
    assert all(v is None or str(v).strip() == "" for v in hsn_values), f"Found non-blank HSNCode values: {[v for v in hsn_values if v is not None and str(v).strip() != ''][:5]}"
    print(f"Test 7 PASSED: All {len(hsn_values)} rows have blank HSNCode as requested.")

    print("\n" + "=" * 80)
    print("ALL 7 SPECIFICATION TESTS PASSED SUCCESSFULLY!")
    print("=" * 80)


if __name__ == "__main__":
    run_tests()
