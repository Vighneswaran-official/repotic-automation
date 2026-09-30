"""
repotic_to_final.py  -  Standalone CLI script for the REPOTIC Automation.

Usage:   python repotic_to_final.py
"""

import sys
import shutil
from pathlib import Path

# Ensure UTF-8 output encoding for Windows consoles
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

# ══════════════════════════════════════════════════════════════════════════════
#  FILE PATHS  -  edit if files move
# ══════════════════════════════════════════════════════════════════════════════
_HERE         = Path(__file__).parent
REPOTIC_PATH  = _HERE / "output" / "upload_repotic.xlsx"
LEDGER_PATH   = _HERE / "output" / "upload_ledger.xlsx"
TEMPLATE_PATH = _HERE / "output" / "upload_template.xlsx"
OUTPUT_PATH   = _HERE / "output" / "Final_Output_Tamilnadu.xlsx"

# ── Constants ──────────────────────────────────────────────────────────────────
TARGET_SECTION = "STATE WISE SALES"
REPOTIC_SHEETS = ["Flipkart", "Meesho", "Snapdeal"]

COL_MAP = {
    "STATE"         : "StateOfSupply",
    "HSN CODE"      : "HSNCode",
    "RATE"          : "TaxPer",
    "QTY"           : "Qty",
    "TAXABLE VALUE" : "TaxableAmt",
    "IGST"          : "IGSTAmt",
    "CGST"          : "CGSTAmt",
    "SGST"          : "SGSTAmt",
    "INVOICE AMOUNT": "Net_Amt",
}

OUTPUT_COLS = [
    "InvNo", "Inv_Dt", "Pty_Name", "Vch_Type", "GSTIN",
    "StateOfSupply", "Product_Name", "HSNCode", "Qty", "UOM",
    "TaxPer", "TaxableAmt", "IGSTAmt", "SGSTAmt", "CGSTAmt",
    "Cess", "OtherAmt", "Net_Amt", "Narration", "Discount",
    "Sales Ledger", "PO No", "PO Date", "DC No", "DC Date",
    "Bill of Supply",
]

NUMERIC_COLS = {"TaxPer", "TaxableAmt", "IGSTAmt", "SGSTAmt", "CGSTAmt", "Net_Amt", "Qty"}

VALIDATION_MAP = {
    "TaxableAmt" : "TAXABLE VALUE",
    "IGSTAmt"    : "IGST",
    "CGSTAmt"    : "CGST",
    "SGSTAmt"    : "SGST",
    "Net_Amt"    : "INVOICE AMOUNT",
}

YELLOW_FILL = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")


# ── Helpers ────────────────────────────────────────────────────────────────────
def _normalize_key(s) -> str:
    return " ".join(str(s).strip().upper().split())


def parse_repotic_sheet(sheet_name: str):
    wb = load_workbook(REPOTIC_PATH, data_only=True)
    if sheet_name not in wb.sheetnames:
        return None, f"Sheet '{sheet_name}' not found"
    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(values_only=True))

    section_idx = None
    for i, row in enumerate(all_rows):
        if row[0] and str(row[0]).strip().upper() == TARGET_SECTION.upper():
            section_idx = i
            break
    if section_idx is None:
        return None, f"'{TARGET_SECTION}' not found in '{sheet_name}'"

    header_idx = None
    for i in range(section_idx + 1, len(all_rows)):
        if any(v is not None for v in all_rows[i]):
            header_idx = i
            break
    if header_idx is None:
        return None, f"No header after '{TARGET_SECTION}' in '{sheet_name}'"

    raw_headers = [
        str(v).strip().upper() if v is not None else f"_COL{j}"
        for j, v in enumerate(all_rows[header_idx])
    ]
    data_rows = []
    for i in range(header_idx + 1, len(all_rows)):
        row = all_rows[i]
        if all(v is None for v in row):
            continue
        if row[0] is None or str(row[0]).strip() == "":
            continue
        data_rows.append(row)

    if not data_rows:
        return None, f"No data rows in '{sheet_name}'"
    return pd.DataFrame(data_rows, columns=raw_headers), None


def load_ledger_all_sheets():
    wb = load_workbook(LEDGER_PATH, data_only=True)
    available = wb.sheetnames
    print(f"[INFO] Ledger: {LEDGER_PATH.name}  |  sheets: {available}")

    ledger = {}
    for sheet_name in available:
        ws = wb[sheet_name]
        header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        col_state = col_sales = col_pty = None
        for idx, h in enumerate(header_row):
            if h is None:
                continue
            hn = str(h).strip().lower()
            if hn == "stateofsupply":
                col_state = idx
            elif hn == "sales ledger":
                col_sales = idx
            elif hn == "pty_name":
                col_pty = idx

        if None in (col_state, col_sales, col_pty):
            print(f"  [WARNING] Sheet '{sheet_name}' missing columns – skipped.")
            continue

        sheet_dict = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            state_raw = row[col_state]
            if state_raw is None or str(state_raw).strip() == "":
                continue
            key = _normalize_key(state_raw)
            sheet_dict[key] = (row[col_pty], row[col_sales])  # exact values

        ledger[sheet_name.strip().lower()] = sheet_dict
        print(f"  Ledger sheet {sheet_name!r}: {len(sheet_dict)} states")

    return ledger, available


def lookup_state(state_raw, sheet_dict: dict):
    if not state_raw or str(state_raw).strip() == "":
        return None, None
    key = _normalize_key(state_raw)
    result = sheet_dict.get(key)
    if result:
        return result  # (Pty_Name, Sales Ledger)
    return None, None


def _to_python(val):
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(val, "item"):
        return val.item()
    return val


# ══════════════════════════════════════════════════════════════════════════════
#  Main
# ══════════════════════════════════════════════════════════════════════════════
def main():
    for p in [REPOTIC_PATH, LEDGER_PATH, TEMPLATE_PATH]:
        if not p.exists():
            print(f"[ERROR] File not found: {p}")
            sys.exit(1)

    OUTPUT_PATH.parent.mkdir(exist_ok=True)
    if OUTPUT_PATH.exists():
        OUTPUT_PATH.unlink()

    ledger, available_sheets = load_ledger_all_sheets()

    all_frames = []
    unmatched  = []

    for sheet_name in REPOTIC_SHEETS:
        print(f"\n{'='*60}")
        print(f"  Sheet: {sheet_name}")
        print(f"{'='*60}")

        norm_sheet = sheet_name.strip().lower()
        if norm_sheet not in ledger:
            print(f"[ERROR] No ledger sheet for '{sheet_name}'. Available: {available_sheets}")
            sys.exit(1)
        sheet_dict = ledger[norm_sheet]

        df_raw, err = parse_repotic_sheet(sheet_name)
        if err:
            print(f"  [SKIP] {err}")
            continue

        n = len(df_raw)
        print(f"  Rows extracted: {n}")

        out_df = pd.DataFrame(index=range(n), columns=OUTPUT_COLS, dtype=object)
        for repotic_col, output_col in COL_MAP.items():
            uc = repotic_col.upper()
            if uc in df_raw.columns:
                out_df[output_col] = df_raw[uc].values
            else:
                print(f"  [WARNING] Column '{repotic_col}' missing – left blank.")

        # Standard format defaults
        out_df["Vch_Type"] = "Auto Sales"
        out_df["Bill of Supply"] = out_df["StateOfSupply"]

        for col in NUMERIC_COLS:
            if col in out_df.columns:
                out_df[col] = pd.to_numeric(out_df[col], errors="coerce").round(2)

        out_df["_marketplace"] = sheet_name

        # Ledger lookup (use .to_dict to preserve "Sales Ledger" key)
        records = out_df.to_dict(orient="records")
        pty_names, sales_ledgs, unmatched_flags = [], [], []

        for rec in records:
            state_val = rec.get("StateOfSupply")
            pty, ledg = lookup_state(state_val, sheet_dict)
            pty_names.append(pty)
            sales_ledgs.append(ledg)
            matched = pty is not None
            if not matched and state_val and str(state_val).strip():
                st = str(state_val).strip()
                print(f"  [UNMATCHED] {sheet_name} – {st} not in ledger sheet '{sheet_name}'")
                unmatched.append({"sheet": sheet_name, "state": st})
            unmatched_flags.append(not matched)

        out_df["Pty_Name"]        = pty_names
        out_df["Sales Ledger"]    = sales_ledgs
        out_df["_unmatched_flag"] = unmatched_flags

        # Validation
        print(f"\n  Validation:")
        for out_col, raw_col in VALIDATION_MAP.items():
            if raw_col in df_raw.columns:
                src = round(pd.to_numeric(df_raw[raw_col], errors="coerce").sum(), 2)
                out = round(pd.to_numeric(out_df[out_col],  errors="coerce").sum(), 2)
                ok  = abs(src - out) <= 0.02
                tag = "OK" if ok else "*** MISMATCH ***"
                print(f"    {out_col:<15} src={src:>13,.2f}  out={out:>13,.2f}  {tag}")

        all_frames.append(out_df)

    if not all_frames:
        print("[ERROR] No data extracted.")
        sys.exit(1)

    combined_df = pd.concat(all_frames, ignore_index=True)
    print(f"\nTotal rows: {len(combined_df)}")

    # Write output
    for col in OUTPUT_COLS:
        if col not in combined_df.columns:
            combined_df[col] = None

    shutil.copy2(TEMPLATE_PATH, OUTPUT_PATH)
    wb_out = load_workbook(OUTPUT_PATH)
    ws_out = wb_out.active
    ws_out.title = "Final Sample Format Tamilnadu"

    header_map = {
        str(ws_out.cell(row=1, column=c).value).strip(): c
        for c in range(1, ws_out.max_column + 1)
        if ws_out.cell(row=1, column=c).value is not None
    }
    pty_col_idx   = header_map.get("Pty_Name")
    sales_col_idx = header_map.get("Sales Ledger")

    data_records = combined_df[OUTPUT_COLS + ["_unmatched_flag"]].to_dict(orient="records")
    for row_offset, record in enumerate(data_records, start=2):
        is_unmatched = bool(record.get("_unmatched_flag", False))
        for col_name in OUTPUT_COLS:
            col_idx = header_map.get(col_name.strip())
            if col_idx is None:
                continue
            val = _to_python(record.get(col_name))
            cell = ws_out.cell(row=row_offset, column=col_idx, value=val)
            if col_name in NUMERIC_COLS and val is not None:
                cell.number_format = "#,##0.00"
        if is_unmatched:
            for cidx in (pty_col_idx, sales_col_idx):
                if cidx:
                    ws_out.cell(row=row_offset, column=cidx).fill = YELLOW_FILL

    wb_out.save(OUTPUT_PATH)

    # Spot-check key rows
    print(f"\n{'='*80}")
    print("SPOT CHECKS (Pty_Name  |  Sales Ledger)")
    print(f"{'='*80}")
    for _, row in combined_df.iterrows():
        state = str(row.get("StateOfSupply", "")).strip().upper()
        mkt   = row.get("_marketplace", "")
        if state in ("TAMIL NADU", "MAHARASHTRA"):
            print(f"  {mkt:<10} {state:<25} | {row.get('Pty_Name')!r}")
            print(f"  {'':<36} | {row.get('Sales Ledger')!r}")

    if unmatched:
        print(f"\nUnmatched states:")
        for u in unmatched:
            print(f"  [{u['sheet']}] {u['state']}")

    print(f"\n[DONE] {len(combined_df)} rows  ->  {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
