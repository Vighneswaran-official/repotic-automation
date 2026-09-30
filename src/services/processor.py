"""
src/services/processor.py - Core data extraction, mapping, validation and Excel generation.
"""

import shutil
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

from src.config import (
    TARGET_SECTION,
    REPOTIC_SHEETS,
    COL_MAP,
    OUTPUT_COLS,
    NUMERIC_COLS,
    VALIDATION_MAP,
    YELLOW_FILL,
    MARKETPLACE_PREFIX_MAP,
    DEFAULT_INVOICE_PREFIX,
    DEFAULT_INVOICE_DATE,
)


def parse_repotic_sheet(repotic_path: Path, sheet_name: str):
    """
    Returns (df_raw, error_string).
    df_raw has uppercase column headers from the STATE WISE SALES section.
    Returns (None, error_msg) on failure.
    """
    wb = load_workbook(repotic_path, data_only=True)
    if sheet_name not in wb.sheetnames:
        return None, f"Sheet '{sheet_name}' not found in {repotic_path.name}"

    ws = wb[sheet_name]
    all_rows = list(ws.iter_rows(values_only=True))

    # Find "STATE WISE SALES" title row
    section_idx = None
    for i, row in enumerate(all_rows):
        if row[0] is not None and str(row[0]).strip().upper() == TARGET_SECTION.upper():
            section_idx = i
            break
    if section_idx is None:
        return None, f"'{TARGET_SECTION}' not found in '{sheet_name}'"

    # Next non-empty row is the column header
    header_idx = None
    for i in range(section_idx + 1, len(all_rows)):
        if any(v is not None for v in all_rows[i]):
            header_idx = i
            break
    if header_idx is None:
        return None, f"No header row after '{TARGET_SECTION}' in '{sheet_name}'"

    raw_headers = [
        str(v).strip().upper() if v is not None else f"_COL{j}"
        for j, v in enumerate(all_rows[header_idx])
    ]

    # Collect data rows
    data_rows = []
    for i in range(header_idx + 1, len(all_rows)):
        row = all_rows[i]
        if all(v is None for v in row):
            continue
        if row[0] is None or str(row[0]).strip() == "":
            continue
        data_rows.append(row)

    if not data_rows:
        return None, f"No data rows under '{TARGET_SECTION}' in '{sheet_name}'"

    return pd.DataFrame(data_rows, columns=raw_headers), None


def _normalize_key(s) -> str:
    """Normalize a state string for lookup: strip, collapse spaces, uppercase."""
    return " ".join(str(s).strip().upper().split())


def load_ledger_all_sheets(ledger_path: Path, logs: list) -> tuple:
    """
    Loads every sheet of the ledger file into a dict:
        { normalized_sheet_name: { normalized_state_key: (Pty_Name, Sales Ledger) } }

    Drops rows where StateOfSupply is empty.
    Logs the number of states found per sheet.
    Values are kept EXACTLY as in the file – no cleanup or title-casing.
    """
    wb = load_workbook(ledger_path, data_only=True)
    available = wb.sheetnames
    logs.append(f"Ledger file: {ledger_path.name}  |  sheets: {available}")

    ledger: dict = {}
    for sheet_name in available:
        ws = wb[sheet_name]
        # Find column indices
        header_row = list(ws.iter_rows(min_row=1, max_row=1, values_only=True))[0]
        col_state = col_sales = col_pty = None
        for idx, h in enumerate(header_row):
            if h is None:
                continue
            hn = str(h).strip().lower()
            if hn == "stateofsupply":
                col_state = idx
            elif hn == "sales ledger":
                col_sales = idx
            elif hn == "pty_name":
                col_pty = idx

        if col_state is None or col_sales is None or col_pty is None:
            logs.append(f"  [WARNING] Sheet '{sheet_name}' missing expected columns – skipped.")
            continue

        sheet_dict: dict = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            state_raw = row[col_state]
            if state_raw is None or str(state_raw).strip() == "":
                continue
            key       = _normalize_key(state_raw)
            pty_val   = row[col_pty]    # keep EXACTLY as-is
            sales_val = row[col_sales]  # keep EXACTLY as-is
            sheet_dict[key] = (pty_val, sales_val)

        norm_sheet = sheet_name.strip().lower()
        ledger[norm_sheet] = sheet_dict
        logs.append(f"  Ledger sheet {sheet_name!r}: {len(sheet_dict)} states")

    return ledger, available


def lookup_state(state_raw, sheet_dict: dict) -> tuple:
    """
    Returns (Pty_Name, Sales Ledger) from the marketplace-specific dict,
    or (None, None) if not found.
    Values are returned EXACTLY as stored.
    """
    if not state_raw or str(state_raw).strip() == "":
        return None, None
    key = _normalize_key(state_raw)
    result = sheet_dict.get(key)
    if result:
        return result  # (Pty_Name, Sales Ledger)
    return None, None


def _to_python(val):
    """Convert numpy/pandas types to standard JSON-safe Python types."""
    if val is None:
        return None
    try:
        if pd.isna(val):
            return None
    except (TypeError, ValueError):
        pass
    if hasattr(val, "item"):
        return val.item()
    return val


def run_automation(repotic_path: Path, ledger_path: Path,
                   template_path: Path, output_path: Path,
                   inv_prefix: str = DEFAULT_INVOICE_PREFIX,
                   inv_date: str = DEFAULT_INVOICE_DATE) -> dict:
    """
    Full pipeline: parse each REPOTIC sheet -> marketplace-specific ledger lookup
    -> write template -> validate sums -> assign sequential InvNo and month-end Inv_Dt.
    Returns a result dict (all values JSON-safe plain Python types).
    """
    logs: list  = []
    unmatched   = []      # list of {sheet, state}
    all_frames  = []
    summary     = []

    # Load ledger
    ledger, available_sheets = load_ledger_all_sheets(ledger_path, logs)

    # Check available sheets in REPOTIC file
    repotic_wb = load_workbook(repotic_path, data_only=True)
    available_repotic = repotic_wb.sheetnames

    # ── Parse + lookup each marketplace sheet ─────────────────────────────────
    for sheet_name in REPOTIC_SHEETS:
        if sheet_name not in available_repotic:
            continue

        logs.append(f"─── Processing: {sheet_name} ───")

        # Verify the ledger has a matching sheet
        norm_sheet = sheet_name.strip().lower()
        if norm_sheet not in ledger:
            raise ValueError(
                f"REPOTIC sheet '{sheet_name}' has no matching ledger sheet. "
                f"Available ledger sheets: {available_sheets}"
            )
        sheet_dict = ledger[norm_sheet]

        df_raw, err = parse_repotic_sheet(repotic_path, sheet_name)
        if err:
            logs.append(f"  [SKIP] {err}")
            continue

        n = len(df_raw)
        logs.append(f"  Rows extracted: {n}")

        # Map columns to output schema
        out_df = pd.DataFrame(index=range(n), columns=OUTPUT_COLS, dtype=object)
        for repotic_col, output_col in COL_MAP.items():
            uc = repotic_col.upper()
            if uc in df_raw.columns:
                out_df[output_col] = df_raw[uc].values
            else:
                logs.append(f"  [WARNING] Column '{repotic_col}' not found – left blank.")

        # Standard format defaults
        out_df["Vch_Type"] = "Auto Sales"
        out_df["Bill of Supply"] = out_df["StateOfSupply"]

        # Generate sequential Invoice Number and Invoice Date
        mkt_code = MARKETPLACE_PREFIX_MAP.get(norm_sheet, sheet_name[:2].upper())
        inv_numbers = [f"{inv_prefix}{mkt_code}-{idx:02d}" for idx in range(1, n + 1)]
        out_df["InvNo"] = inv_numbers
        out_df["Inv_Dt"] = inv_date

        # Coerce numeric columns
        for col in NUMERIC_COLS:
            if col in out_df.columns:
                out_df[col] = pd.to_numeric(out_df[col], errors="coerce").round(2)

        # Add marketplace tag for preview
        out_df["_marketplace"] = sheet_name

        # Ledger lookup – use .to_dict() to preserve "Sales Ledger" key exactly
        records = out_df.to_dict(orient="records")
        pty_names, sales_ledgs, unmatched_flags = [], [], []

        for rec in records:
            state_val = rec.get("StateOfSupply")
            pty, ledg = lookup_state(state_val, sheet_dict)
            pty_names.append(pty)
            sales_ledgs.append(ledg)
            matched = pty is not None

            if not matched and state_val and str(state_val).strip():
                st = str(state_val).strip()
                logs.append(f"  [UNMATCHED] {sheet_name} – {st} not in ledger sheet '{sheet_name}'")
                unmatched.append({"sheet": sheet_name, "state": st})
            unmatched_flags.append(not matched)

        out_df["Pty_Name"]       = pty_names
        out_df["Sales Ledger"]   = sales_ledgs
        out_df["_unmatched_flag"]= unmatched_flags

        # Validation: compare output sums with source sums
        sheet_ok = True
        val_detail = []
        for out_col, raw_col in VALIDATION_MAP.items():
            if raw_col in df_raw.columns:
                src = float(round(pd.to_numeric(df_raw[raw_col], errors="coerce").sum(), 2))
                out = float(round(pd.to_numeric(out_df[out_col],  errors="coerce").sum(), 2))
                ok  = bool(abs(src - out) <= 0.02)
                if not ok:
                    sheet_ok = False
                val_detail.append({"col": out_col, "source": src, "output": out, "ok": ok})
                tag = "OK" if ok else "MISMATCH"
                logs.append(f"    {out_col:<15} src={src:>13,.2f}  out={out:>13,.2f}  {tag}")

        summary.append({
            "sheet"      : sheet_name,
            "rows"       : int(n),
            "TaxableAmt" : float(round(pd.to_numeric(out_df["TaxableAmt"], errors="coerce").sum(), 2)),
            "IGSTAmt"    : float(round(pd.to_numeric(out_df["IGSTAmt"],    errors="coerce").sum(), 2)),
            "CGSTAmt"    : float(round(pd.to_numeric(out_df["CGSTAmt"],    errors="coerce").sum(), 2)),
            "SGSTAmt"    : float(round(pd.to_numeric(out_df["SGSTAmt"],    errors="coerce").sum(), 2)),
            "Net_Amt"    : float(round(pd.to_numeric(out_df["Net_Amt"],    errors="coerce").sum(), 2)),
            "status"     : "OK" if sheet_ok else "MISMATCH",
            "validation" : val_detail,
        })

        all_frames.append(out_df)

    if not all_frames:
        raise ValueError("No data extracted from any sheet. Check the input file.")

    combined_df = pd.concat(all_frames, ignore_index=True)
    logs.append(f"Total rows written: {int(len(combined_df))}")

    # ── Write output workbook ─────────────────────────────────────────────────
    shutil.copy2(template_path, output_path)
    wb_out = load_workbook(output_path)
    ws_out = wb_out.active
    ws_out.title = "Final Sample Format Tamilnadu"

    # Build header map (strip to avoid invisible-space mismatches)
    header_map = {
        str(ws_out.cell(row=1, column=c).value).strip(): c
        for c in range(1, ws_out.max_column + 1)
        if ws_out.cell(row=1, column=c).value is not None
    }

    # Column indices for cells to highlight yellow when unmatched
    pty_col_idx   = header_map.get("Pty_Name")
    sales_col_idx = header_map.get("Sales Ledger")

    # Use .to_dict() to preserve "Sales Ledger" key exactly
    data_records = combined_df[OUTPUT_COLS + ["_unmatched_flag"]].to_dict(orient="records")
    for row_offset, record in enumerate(data_records, start=2):
        is_unmatched = bool(record.get("_unmatched_flag", False))
        for col_name in OUTPUT_COLS:
            col_idx = header_map.get(col_name.strip())
            if col_idx is None:
                continue
            val = _to_python(record.get(col_name))
            cell = ws_out.cell(row=row_offset, column=col_idx, value=val)
            if col_name in NUMERIC_COLS and val is not None:
                cell.number_format = "#,##0.00"
        # Highlight Pty_Name and Sales Ledger yellow for unmatched rows
        if is_unmatched:
            for cidx in (pty_col_idx, sales_col_idx):
                if cidx:
                    ws_out.cell(row=row_offset, column=cidx).fill = YELLOW_FILL

    wb_out.save(output_path)
    logs.append(f"Saved: {output_path.name}  (sheet: '{ws_out.title}')")

    # ── Preview (first 10 rows, JSON-safe) ───────────────────────────────────
    preview_cols = [
        "InvNo", "Inv_Dt", "_marketplace", "Vch_Type", "StateOfSupply", "Bill of Supply",
        "HSNCode", "Qty", "TaxPer", "TaxableAmt", "IGSTAmt", "SGSTAmt", "CGSTAmt",
        "Net_Amt", "Pty_Name", "Sales Ledger",
    ]
    preview_records = [
        {col: _to_python(row[col]) for col in preview_cols}
        for _, row in combined_df[preview_cols].head(10).iterrows()
    ]

    return {
        "logs"      : logs,
        "summary"   : summary,
        "unmatched" : unmatched,
        "preview"   : preview_records,
        "total_rows": int(len(combined_df)),
    }
